"""Pruebas del reporte de consumo por usuario y por zona del administrador (PB032)."""
from datetime import date, datetime, timedelta
from io import BytesIO
from unittest.mock import patch

from django.core.cache import cache, caches
from django.http import QueryDict
from django.test import TestCase, override_settings
from django.utils import timezone
from openpyxl import load_workbook

from ml_engine.models import ConsumptionAggregate, Home, Zone

from . import consumo_admin
from .models import AjusteComunidad, Rol, Usuario

HOY = date(2026, 10, 20)

VIVIENDAS = [
    {'id_vivienda': 1, 'nic': 'NIC-1', 'direccion': 'Calle 1', 'id_usuario_propietario': 11, 'nombre_completo_titular': 'Ana'},
    {'id_vivienda': 2, 'nic': 'NIC-2', 'direccion': 'Calle 2', 'id_usuario_propietario': 12, 'nombre_completo_titular': 'Beto'},
    {'id_vivienda': 3, 'nic': 'NIC-3', 'direccion': 'Calle 3', 'id_usuario_propietario': None, 'nombre_completo_titular': 'Carla Titular'},
]
DUENOS = {
    11: {'id_usuario': 11, 'nombre': 'Ana', 'apellido': 'López', 'correo': 'ana@example.com'},
    12: {'id_usuario': 12, 'nombre': 'Beto', 'apellido': 'Ruiz', 'correo': 'beto@example.com'},
}


def _hora(dia, hora=12):
    return timezone.make_aware(datetime(dia.year, dia.month, dia.day, hora))


def _select_falso(registrado):
    """Supabase de mentira: viviendas y la tabla `consumo`."""
    def select(tabla, columnas='*', params=None):
        if tabla == 'vivienda':
            filtro = (params or {}).get('id_vivienda', '')
            if filtro.startswith('eq.'):
                return [v for v in VIVIENDAS if str(v['id_vivienda']) == filtro[3:]]
            return VIVIENDAS
        if tabla == 'consumo':
            return registrado
        return []
    return select


@override_settings(CONSUMO_LIMITE_M3_MES=3, REVALIDAR_SESION_SEGUNDOS=0)
class ReporteConsumoTests(TestCase):
    databases = {'default', 'timeseries'}

    def setUp(self):
        cache.clear()
        caches['limites'].clear()
        zona_norte = Zone.objects.create(name='Tanque 1', source='real')
        zona_sur = Zone.objects.create(name='Tanque 2', source='real')
        self.hogares = {
            1: Home.objects.create(zone=zona_norte, source='real', source_ref_vivienda_id=1, meta={'zona': 'Zona Norte'}),
            2: Home.objects.create(zone=zona_norte, source='real', source_ref_vivienda_id=2, meta={'zona': 'Zona Norte'}),
            3: Home.objects.create(zone=zona_sur, source='real', source_ref_vivienda_id=3, meta={'zona': 'Zona Sur'}),
        }

    def _medir(self, vivienda, dia, litros):
        ConsumptionAggregate.objects.create(home=self.hogares[vivienda], period='hourly', period_start=_hora(dia),
                                            litros=litros)

    def _reporte(self, registrado=(), desde=date(2026, 10, 1), hasta=HOY):
        with patch('App.consumo_admin.supabase_client.select', side_effect=_select_falso(list(registrado))), \
                patch('App.consumo_admin.users_by_id', return_value=DUENOS):
            return consumo_admin.reporte(desde, hasta)

    def test_usa_lo_registrado_y_si_no_hay_lo_medido_por_el_flujo_en_m3(self):
        self._medir(1, date(2026, 10, 5), 500)       # medido por el flujo
        self._medir(1, date(2026, 10, 6), 999)       # este día hay registro: gana el registro
        registrado = [{'id_vivienda': 1, 'fecha': _hora(date(2026, 10, 6)).isoformat(), 'consumo_total': 250}]
        datos = self._reporte(registrado)
        ana = next(u for u in datos['usuarios'] if u['id_vivienda'] == 1)
        self.assertEqual(ana['m3'], 0.75)            # 500 L + 250 L = 0.75 m³
        self.assertEqual(ana['usuario'], 'Ana López')
        self.assertEqual(ana['zona'], 'Zona Norte')
        carla = next(u for u in datos['usuarios'] if u['id_vivienda'] == 3)
        self.assertEqual(carla['usuario'], 'Carla Titular')   # sin cuenta vinculada: el titular
        self.assertFalse(carla['vinculada'])

    def test_resalta_a_quien_pasa_el_limite_y_la_zona_afectada(self):
        # Límite 3 m³ al mes -> 2 m³ para 20 días. Beto gasta 2.5 m³.
        self._medir(2, date(2026, 10, 10), 2500)
        self._medir(1, date(2026, 10, 10), 100)
        datos = self._reporte()
        self.assertEqual(datos['limite_periodo'], 2.0)
        beto = next(u for u in datos['usuarios'] if u['id_vivienda'] == 2)
        self.assertTrue(beto['excede'])
        self.assertEqual(beto['exceso_m3'], 0.5)
        self.assertEqual(datos['usuarios'][0]['id_vivienda'], 2)   # el que más gasta va primero
        norte = next(z for z in datos['zonas'] if z['zona'] == 'Zona Norte')
        sur = next(z for z in datos['zonas'] if z['zona'] == 'Zona Sur')
        self.assertTrue(norte['afectada'])
        self.assertFalse(sur['afectada'])
        self.assertEqual(norte['viviendas'], 2)
        self.assertTrue(any('pasan el límite' in frase for frase in datos['hallazgos']))
        self.assertTrue(any('Zona Norte' in frase for frase in datos['hallazgos']))

    def test_aumento_anormal_frente_al_periodo_anterior(self):
        desde, hasta = date(2026, 10, 11), date(2026, 10, 20)
        self._medir(3, date(2026, 10, 5), 200)      # periodo anterior (1 al 10)
        self._medir(3, date(2026, 10, 15), 400)     # este periodo: el doble
        datos = self._reporte(desde=desde, hasta=hasta)
        carla = next(u for u in datos['usuarios'] if u['id_vivienda'] == 3)
        self.assertEqual(carla['m3_anterior'], 0.2)
        self.assertEqual(carla['cambio'], 1.0)
        self.assertTrue(carla['anormal'])
        self.assertEqual(datos['anterior_desde'], date(2026, 10, 1))

    def test_solo_cuenta_el_periodo_elegido(self):
        self._medir(1, date(2026, 9, 30), 5000)
        self._medir(1, date(2026, 10, 2), 300)
        datos = self._reporte()
        self.assertEqual(next(u for u in datos['usuarios'] if u['id_vivienda'] == 1)['m3'], 0.3)
        self.assertEqual(len(datos['tendencia']), 20)

    def test_periodo_por_mes_por_rango_o_el_mes_en_curso(self):
        self.assertEqual(consumo_admin.periodo_desde_parametros(QueryDict('mes=2026-09'), HOY),
                         (date(2026, 9, 1), date(2026, 9, 30), 'Septiembre 2026'))
        self.assertEqual(consumo_admin.periodo_desde_parametros(QueryDict('desde=2026-10-15&hasta=2026-10-05'), HOY)[:2],
                         (date(2026, 10, 5), date(2026, 10, 15)))
        self.assertEqual(consumo_admin.periodo_desde_parametros(QueryDict('mes=basura'), HOY)[:2],
                         (date(2026, 10, 1), HOY))
        # El mes en curso llega hasta hoy, no al último día del mes.
        self.assertEqual(consumo_admin.periodo_desde_parametros(QueryDict('mes=2026-10'), HOY)[1], HOY)

    def test_el_limite_guardado_manda_sobre_el_de_la_configuracion(self):
        self.assertEqual(consumo_admin.limite_mensual(), 3)
        consumo_admin.guardar_limite(12.5, 99)
        self.assertEqual(consumo_admin.limite_mensual(), 12.5)
        self.assertEqual(AjusteComunidad.objects.get().id_usuario, 99)


@override_settings(CONSUMO_LIMITE_M3_MES=3, REVALIDAR_SESION_SEGUNDOS=0)
class PantallaConsumoAdminTests(TestCase):
    databases = {'default', 'timeseries'}

    def setUp(self):
        cache.clear()
        caches['limites'].clear()
        rol_admin = Rol.objects.create(id_rol=2, nombre_rol='admin')
        rol_usuario = Rol.objects.create(id_rol=1, nombre_rol='user')
        self.admin = Usuario.objects.create_user(email='admin@example.com', nombre='Ad', apellido='Min',
                                                 password='clave-segura-123', rol=rol_admin, supabase_id=99)
        self.vecino = Usuario.objects.create_user(email='vecino@example.com', nombre='Ve', apellido='Cino',
                                                  password='clave-segura-123', rol=rol_usuario, supabase_id=11)
        zona = Zone.objects.create(name='Tanque 1', source='real')
        hogar = Home.objects.create(zone=zona, source='real', source_ref_vivienda_id=2, meta={'zona': 'Zona Norte'})
        ConsumptionAggregate.objects.create(home=hogar, period='hourly', period_start=timezone.now() - timedelta(hours=1),
                                            litros=9000)
        self.parches = [
            patch('App.consumo_admin.supabase_client.select', side_effect=_select_falso([])),
            patch('App.consumo_admin.users_by_id', return_value=DUENOS),
        ]
        for parche in self.parches:
            parche.start()
            self.addCleanup(parche.stop)

    def test_solo_el_admin_ve_el_reporte(self):
        self.assertEqual(self.client.get('/admin-consumo/').status_code, 302)          # sin sesión
        self.client.force_login(self.vecino)
        self.assertRedirects(self.client.get('/admin-consumo/'), '/', fetch_redirect_response=False)
        self.assertRedirects(self.client.get('/admin-consumo/exportar/'), '/', fetch_redirect_response=False)

    def test_el_admin_ve_usuarios_zonas_y_el_resaltado(self):
        self.client.force_login(self.admin)
        respuesta = self.client.get('/admin-consumo/')
        self.assertEqual(respuesta.status_code, 200)
        html = respuesta.content.decode()
        self.assertIn('Reporte de consumo por usuario', html)
        self.assertIn('Beto Ruiz', html)
        self.assertIn('Zona Norte', html)
        self.assertIn('fila-alerta', html)               # Beto gasta 9 m³ y el límite es 3 al mes
        self.assertIn('Pasa el límite por', html)
        self.assertIn('Periodo de facturación', html)
        # Filtrar por zona deja solo esa zona en la tabla de usuarios.
        solo_sur = self.client.get('/admin-consumo/?zona=Sin+zona').content.decode()
        self.assertNotIn('Beto Ruiz</th>', solo_sur.split('Reporte de consumo por usuario')[1].split('<caption>')[0])

    def test_el_admin_cambia_el_limite(self):
        self.client.force_login(self.admin)
        self.client.post('/admin-consumo/?mes=2026-10', {'limite': '15'})
        self.assertEqual(AjusteComunidad.objects.get().valor, 15)
        self.client.post('/admin-consumo/', {'limite': '-4'})
        self.assertEqual(AjusteComunidad.objects.get().valor, 15)   # un valor inválido no se guarda

    def test_el_panel_general_muestra_el_consumo_de_cada_usuario(self):
        self.client.force_login(self.admin)
        with patch('App.views.run_parallel', return_value=([], [], [], [])):
            html = self.client.get('/admin-panel/').content.decode()
        self.assertIn('Consumo por usuario', html)
        self.assertIn('Beto Ruiz', html)
        self.assertIn('/admin-consumo/vivienda/2/', html)
        self.assertIn('Pasa el límite', html)

    def test_detalle_de_un_usuario_dia_por_dia(self):
        self.client.force_login(self.admin)
        respuesta = self.client.get('/admin-consumo/vivienda/2/')
        self.assertEqual(respuesta.status_code, 200)
        html = respuesta.content.decode()
        self.assertIn('Beto Ruiz', html)
        self.assertIn('beto@example.com', html)
        self.assertIn('Detalle por día', html)
        self.assertRegex(html, r'(Lunes|Martes|Miércoles|Jueves|Viernes|Sábado|Domingo) \d\d/\d\d/\d{4}')
        self.assertIn('Sobre el límite diario', html)       # 9 m³ en un día, el límite es 0.1 al día
        self.assertIn('Este usuario pasa el límite de la comunidad', html)
        # Una vivienda que no existe regresa al reporte.
        self.assertRedirects(self.client.get('/admin-consumo/vivienda/999/'), '/admin-consumo/',
                             fetch_redirect_response=False)
        self.client.force_login(self.vecino)
        self.assertRedirects(self.client.get('/admin-consumo/vivienda/2/'), '/', fetch_redirect_response=False)

    def test_descarga_el_excel_con_resumen_usuarios_zonas_y_dias(self):
        self.client.force_login(self.admin)
        respuesta = self.client.get('/admin-consumo/exportar/')
        self.assertEqual(respuesta.status_code, 200)
        self.assertIn('spreadsheetml', respuesta['Content-Type'])
        self.assertIn('attachment; filename="consumo_smartdrop_', respuesta['Content-Disposition'])
        libro = load_workbook(BytesIO(respuesta.content))
        self.assertEqual(libro.sheetnames, ['Resumen', 'Por usuario', 'Por zona', 'Por día'])
        hoja = libro['Por usuario']
        self.assertEqual(hoja['A1'].value, 'Usuario')
        self.assertEqual(hoja['G1'].value, 'Consumo (m³)')
        fila_beto = next(fila for fila in hoja.iter_rows(min_row=2, values_only=True) if fila[0] == 'Beto Ruiz')
        self.assertEqual(fila_beto[6], 9.0)
        self.assertEqual(fila_beto[10], 'Sí')
