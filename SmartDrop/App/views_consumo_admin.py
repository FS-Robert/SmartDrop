"""Pantalla "Consumo por usuario" del administrador (PB032) y su descarga en Excel.

/admin-consumo/                 reporte en pantalla (por usuario y por zona, con filtros de periodo)
/admin-consumo/exportar/        el mismo reporte en Excel (.xlsx)
El cálculo está en consumo_admin.py; aquí solo se arma la página y el archivo.
"""
import logging
from io import BytesIO

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.shortcuts import redirect, render
from django.urls import reverse
from django.views.decorators.http import require_http_methods
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import consumo_admin, supabase_client
from .queries import is_admin, owner_id
from .seguridad import celda_csv, limitar

logger = logging.getLogger(__name__)

LIMITE_MINIMO_M3 = 1
LIMITE_MAXIMO_M3 = 500


def _datos(request):
    desde, hasta, etiqueta = consumo_admin.periodo_desde_parametros(request.GET)
    datos = consumo_admin.reporte(desde, hasta)
    datos['etiqueta'] = etiqueta
    return datos


def _filtro_zona(datos, zona):
    if zona and any(z['zona'] == zona for z in datos['zonas']):
        return [u for u in datos['usuarios'] if u['zona'] == zona]
    return datos['usuarios']


@login_required(login_url='login')
@require_http_methods(['GET', 'POST'])
def admin_consumo(request):
    if not is_admin(request.user):
        return redirect('dashboard')

    if request.method == 'POST':
        # El administrador cambia el límite de consumo de la comunidad.
        try:
            valor = float(str(request.POST.get('limite', '')).replace(',', '.'))
        except ValueError:
            valor = None
        if valor is None or not LIMITE_MINIMO_M3 <= valor <= LIMITE_MAXIMO_M3:
            messages.error(request, f'El límite debe ser un número entre {LIMITE_MINIMO_M3} y {LIMITE_MAXIMO_M3} m³.')
        else:
            consumo_admin.guardar_limite(round(valor, 2), owner_id(request.user))
            messages.success(request, f'Listo: el límite de la comunidad ahora es {valor:g} m³ al mes por vivienda.')
        return redirect(f"{reverse('admin_consumo')}?{request.GET.urlencode()}")

    try:
        datos = _datos(request)
    except supabase_client.SupabaseError:
        logger.exception('No se pudo armar el reporte de consumo')
        messages.error(request, 'No se pudieron cargar las viviendas. Inténtalo de nuevo en un momento.')
        datos = None

    zona = request.GET.get('zona', '')
    hay_rango = request.GET.get('desde') and request.GET.get('hasta')
    contexto = {
        'datos': datos,
        'usuarios': _filtro_zona(datos, zona) if datos else [],
        'zona_elegida': zona,
        'mes': '' if hay_rango else (request.GET.get('mes') or (datos['desde'].strftime('%Y-%m') if datos else '')),
        'desde_txt': request.GET.get('desde', ''),
        'hasta_txt': request.GET.get('hasta', ''),
        'consulta': request.GET.urlencode(),
        'aumento_pct': round(consumo_admin.AUMENTO_ANORMAL * 100),
    }
    if datos:
        contexto['grafica_zonas'] = {
            'etiquetas': [z['zona'] for z in datos['zonas']],
            'm3': [z['m3'] for z in datos['zonas']],
            'afectadas': [z['afectada'] for z in datos['zonas']],
        }
        contexto['grafica_dias'] = {
            'etiquetas': [d['etiqueta'] for d in datos['tendencia']],
            'm3': [d['m3'] for d in datos['tendencia']],
            'limite_dia': round(datos['limite_mes'] / consumo_admin.DIAS_DEL_MES * len(datos['usuarios']), 3),
        }
    return render(request, 'App/admin_consumo.html', contexto)


# ── Excel ────────────────────────────────────────────────────────────────────

_TITULO = Font(bold=True, size=14, color='0B3D5C')
_NEGRITA = Font(bold=True, color='FFFFFF')
_CABECERA = PatternFill('solid', fgColor='0369A1')
_EXCEDE = PatternFill('solid', fgColor='FDE2E1')
_ANORMAL = PatternFill('solid', fgColor='FEF3C7')


def _hoja(libro, titulo, encabezado, filas, anchos):
    hoja = libro.create_sheet(titulo)
    hoja.append(encabezado)
    for celda in hoja[1]:
        celda.font, celda.fill = _NEGRITA, _CABECERA
        celda.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for fila in filas:
        hoja.append([celda_csv(v) for v in fila])
    for indice, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(indice)].width = ancho
    hoja.freeze_panes = 'A2'
    return hoja


def libro_excel(datos, usuarios):
    libro = Workbook()
    resumen = libro.active
    resumen.title = 'Resumen'
    resumen['A1'] = 'SmartDrop — Reporte de consumo por usuario'
    resumen['A1'].font = _TITULO
    filas = [
        ('Periodo de facturación', datos['etiqueta']),
        ('Desde', datos['desde'].strftime('%d/%m/%Y')),
        ('Hasta', datos['hasta'].strftime('%d/%m/%Y')),
        ('Días', datos['dias']),
        ('Consumo total (m³)', datos['total_m3']),
        ('Promedio por vivienda (m³)', datos['promedio_m3']),
        ('Límite de la comunidad (m³ al mes)', datos['limite_mes']),
        ('Límite para este periodo (m³)', datos['limite_periodo']),
        ('Viviendas que pasan el límite', datos['excedidos']),
        ('Viviendas con aumento anormal', datos['anormales']),
        ('Datos medidos hasta', datos['ultima_lectura'].strftime('%d/%m/%Y %H:%M') if datos['ultima_lectura'] else '—'),
    ]
    for fila, (nombre, valor) in enumerate(filas, start=3):
        resumen.cell(row=fila, column=1, value=nombre).font = Font(bold=True)
        resumen.cell(row=fila, column=2, value=valor)
    fila = len(filas) + 4
    resumen.cell(row=fila, column=1, value='Hallazgos').font = Font(bold=True)
    for numero, frase in enumerate(datos['hallazgos'], start=1):
        resumen.cell(row=fila + numero, column=1, value=celda_csv(f'• {frase}'))
    resumen.column_dimensions['A'].width = 38
    resumen.column_dimensions['B'].width = 28

    hoja = _hoja(
        libro, 'Por usuario',
        ['Usuario', 'Correo', 'NIC', 'Dirección', 'Zona', 'Periodo de facturación', 'Consumo (m³)',
         'Promedio diario (m³)', 'Periodo anterior (m³)', 'Cambio (%)', 'Pasa el límite', 'Exceso (m³)',
         'Aumento anormal'],
        [[u['usuario'], u['correo'], u['nic'], u['direccion'], u['zona'], datos['etiqueta'], u['m3'], u['m3_dia'],
          u['m3_anterior'], round(u['cambio'] * 100, 1) if u['cambio'] is not None else '',
          'Sí' if u['excede'] else 'No', u['exceso_m3'], 'Sí' if u['anormal'] else 'No'] for u in usuarios],
        [28, 28, 12, 34, 16, 22, 13, 15, 16, 11, 13, 11, 15],
    )
    for numero, usuario in enumerate(usuarios, start=2):
        relleno = _EXCEDE if usuario['excede'] else _ANORMAL if usuario['anormal'] else None
        if relleno:
            for celda in hoja[numero]:
                celda.fill = relleno

    _hoja(
        libro, 'Por zona',
        ['Zona', 'Viviendas', 'Consumo (m³)', 'Promedio por vivienda (m³)', '% del total', 'Cambio (%)',
         'Pasan el límite', 'Aumento anormal', 'Revisar'],
        [[z['zona'], z['viviendas'], z['m3'], z['promedio_m3'], z['porcentaje'],
          round(z['cambio'] * 100, 1) if z['cambio'] is not None else '', z['excedidos'], z['anormales'],
          'Sí' if z['afectada'] else 'No'] for z in datos['zonas']],
        [22, 11, 14, 22, 11, 11, 14, 15, 9],
    )
    _hoja(libro, 'Por día', ['Día', 'Consumo de todas las viviendas (m³)'],
          [[d['dia'], d['m3']] for d in datos['tendencia']], [14, 32])

    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


@login_required(login_url='login')
@limitar('exportar', por='usuario', metodos=('GET',))
def admin_consumo_exportar(request):
    if not is_admin(request.user):
        return redirect('dashboard')
    try:
        datos = _datos(request)
    except supabase_client.SupabaseError:
        messages.error(request, 'No se pudo generar el archivo. Inténtalo de nuevo en un momento.')
        return redirect('admin_consumo')
    usuarios = _filtro_zona(datos, request.GET.get('zona', ''))
    respuesta = HttpResponse(
        libro_excel(datos, usuarios),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    nombre = f"consumo_smartdrop_{datos['desde']:%Y%m%d}_{datos['hasta']:%Y%m%d}.xlsx"
    respuesta['Content-Disposition'] = f'attachment; filename="{nombre}"'
    return respuesta


@login_required(login_url='login')
def admin_consumo_usuario(request, id_vivienda):
    """Consumo de un usuario (su vivienda) día por día, con su gráfica."""
    if not is_admin(request.user):
        return redirect('dashboard')
    desde, hasta, etiqueta = consumo_admin.periodo_desde_parametros(request.GET)
    try:
        datos = consumo_admin.detalle(id_vivienda, desde, hasta)
    except supabase_client.SupabaseError:
        messages.error(request, 'No se pudo cargar el consumo de esta vivienda. Inténtalo de nuevo en un momento.')
        return redirect('admin_consumo')
    if datos is None:
        messages.error(request, 'Esa vivienda no existe.')
        return redirect('admin_consumo')
    datos['etiqueta'] = etiqueta
    hay_rango = request.GET.get('desde') and request.GET.get('hasta')
    return render(request, 'App/admin_consumo_usuario.html', {
        'datos': datos,
        'mes': '' if hay_rango else (request.GET.get('mes') or desde.strftime('%Y-%m')),
        'desde_txt': request.GET.get('desde', ''),
        'hasta_txt': request.GET.get('hasta', ''),
        'consulta': request.GET.urlencode(),
        'grafica': {
            'etiquetas': [f"{item['dia']:%d/%m}" for item in datos['dias']],
            'm3': [item['m3'] for item in datos['dias']],
            'limite_dia': datos['limite_dia'],
        },
    })
